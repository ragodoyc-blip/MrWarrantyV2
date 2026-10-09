import os
import re
import shutil
import subprocess
import time
from pathlib import Path

import requests
from pdf2image import convert_from_path
from pdf2image.exceptions import PDFPageCountError

try:
    from PIL import Image, ImageDraw, ImageFont
except ImportError:
    Image = None  # type: ignore

from azure.storage.blob import (
    BlobServiceClient,
    BlobSasPermissions,
    generate_blob_sas,
)
from datetime import datetime, timedelta

from mr_warranty.config.config import (
    ADJUNTOS_DIR,
    AZURE_STORAGE_CONNECTION_STRING,
    AZURE_STORAGE_CONTAINER,
    POPPLER_PATH,
)
from mr_warranty.core.logger import log
from mr_warranty.services.mintral import Ejecutar_Mistral_diccionario
from mr_warranty.adapters.sqis_client import URL_Adjuntos


# ── Azure Blob Storage ──────────────────────────────────────
_blob_service_client = BlobServiceClient.from_connection_string(AZURE_STORAGE_CONNECTION_STRING)
_container_client = _blob_service_client.get_container_client(AZURE_STORAGE_CONTAINER)


def generar_sas(blob_name: str) -> str:
    sas_token = generate_blob_sas(
        account_name=_blob_service_client.account_name,
        container_name=AZURE_STORAGE_CONTAINER,
        blob_name=blob_name,
        account_key=_blob_service_client.credential.account_key,
        permission=BlobSasPermissions(read=True),
        expiry=datetime.utcnow() + timedelta(minutes=30),
    )
    blob_client = _container_client.get_blob_client(blob_name)
    return f"{blob_client.url}?{sas_token}"


def obtener_urls_blob_existentes(
    blob_folder: str,
    filename_prefix: str | None = None,
) -> list[str]:
    """Retorna SAS URLs de imágenes existentes sin volver a subirlas."""
    prefix = f"{blob_folder.strip('/')}/"
    try:
        blob_names = []
        for blob in _container_client.list_blobs(name_starts_with=prefix):
            name = blob.name
            filename = name.rsplit("/", 1)[-1]
            if not filename.lower().endswith((".png", ".jpg", ".jpeg")):
                continue
            if filename_prefix and not filename.startswith(filename_prefix):
                continue
            blob_names.append(name)

        return [generar_sas(name) for name in sorted(blob_names)]
    except Exception as e:
        log.warning("No se pudieron consultar blobs existentes en %s: %s", blob_folder, e)
        return []


def subir_imagenes_blob(document_folder: str, local_folder: str) -> list[str]:
    """Sube todas las imágenes al Blob y retorna URLs SAS válidas por 30 minutos."""
    urls = []
    for filename in sorted(os.listdir(local_folder)):
        if filename.lower().endswith((".png", ".jpg", ".jpeg")):
            local_path = os.path.join(local_folder, filename)
            blob_name = f"{document_folder}/{filename}"
            blob_client = _container_client.get_blob_client(blob_name)

            with open(local_path, "rb") as data:
                blob_client.upload_blob(data, overwrite=True)

            sas_url = generar_sas(blob_name)
            urls.append(sas_url)
    return urls


def reconstruir_urls_existentes(download_path: str) -> dict[str, list[str]]:
    documentos_urls = {}
    # Path.name funciona con rutas Windows y evita incluir "AdjuntosSQIS\\" en el Blob.
    id_claim = Path(download_path).name

    for carpeta in sorted(os.listdir(download_path)):
        carpeta_path = os.path.join(download_path, carpeta)
        if not os.path.isdir(carpeta_path):
            continue

        urls = []
        for filename in sorted(os.listdir(carpeta_path)):
            if filename.lower().endswith((".png", ".jpg", ".jpeg")):
                blob_name = f"{id_claim}/{carpeta}/{filename}"
                sas_url = generar_sas(blob_name)
                urls.append(sas_url)

        documentos_urls[carpeta] = urls
    return documentos_urls


def cerrar_edge() -> None:
    try:
        subprocess.run("taskkill /F /IM msedge.exe /T", shell=True)
        subprocess.run("taskkill /F /IM msedgedriver.exe /T", shell=True)
        time.sleep(2)
    except Exception as e:
        log.warning("No se pudieron cerrar procesos de Edge: %s", e)


def pdf_to_imagenes(pdf_file_path: str, output_folder: str) -> list[str]:
    if not os.path.exists(pdf_file_path):
        raise FileNotFoundError(f"El archivo PDF no existe: {pdf_file_path}")

    try:
        log.info("Procesando archivo PDF: %s", pdf_file_path)
        pages = convert_from_path(pdf_file_path, dpi=200, poppler_path=POPPLER_PATH)
        generated_files = []
        base_name = os.path.splitext(os.path.basename(pdf_file_path))[0]

        for idx, page in enumerate(pages, start=1):
            out_path = os.path.join(output_folder, f"{base_name}_page_{idx:02}.png")
            page.save(out_path, "PNG")
            generated_files.append(out_path)

        return generated_files

    except PDFPageCountError as e:
        log.error("Error al contar páginas del PDF %s: %s", pdf_file_path, e)
        raise ValueError(f"No se pudo procesar el archivo PDF: {pdf_file_path}. Error: {e}")
    except Exception as e:
        log.error("Error inesperado al procesar PDF %s: %s", pdf_file_path, e)
        raise ValueError(f"Error inesperado al procesar el archivo PDF: {pdf_file_path}. Error: {e}")


def excel_to_imagenes(excel_file_path: str, output_folder: str) -> list[str]:
    """Convierte un Excel (PLM) a imagen para IA: extrae texto y genera PNG."""
    if not os.path.exists(excel_file_path):
        raise FileNotFoundError(f"El archivo Excel no existe: {excel_file_path}")
    if Image is None:
        raise ValueError("Pillow no disponible para convertir Excel a imagen")

    try:
        log.info("Procesando archivo Excel PLM: %s", excel_file_path)
        # Extrae texto con openpyxl/pandas
        text_lines: list[str] = []
        try:
            import openpyxl

            wb = openpyxl.load_workbook(excel_file_path, read_only=True, data_only=True)
            for ws in wb.worksheets:
                text_lines.append(f"Hoja: {ws.title}")
                for row in ws.iter_rows(values_only=True):
                    if row and any(v is not None and str(v).strip() for v in row):
                        vals = [str(v).strip() for v in row if v is not None and str(v).strip()]
                        text_lines.append(" | ".join(vals)[:200])
                        if len(text_lines) > 60:
                            break
                if len(text_lines) > 60:
                    break
        except Exception:
            import pandas as pd

            xls = pd.ExcelFile(excel_file_path)
            for sheet in xls.sheet_names[:2]:
                df = xls.parse(sheet, nrows=20)
                text_lines.append(f"Hoja: {sheet}")
                text_lines.append(df.to_string()[:2000])

        if not text_lines:
            text_lines = ["PLM Excel sin contenido legible"]

        # Genera imagen con texto
        base_name = os.path.splitext(os.path.basename(excel_file_path))[0]
        out_path = os.path.join(output_folder, f"{base_name}_sheet_01.png")
        # Imagen 1200x800
        img = Image.new("RGB", (1200, 800), "white")
        draw = ImageDraw.Draw(img)
        try:
            font = ImageFont.truetype("arial.ttf", 14)
        except Exception:
            font = ImageFont.load_default()
        y = 10
        for line in text_lines[:35]:
            draw.text((10, y), line[:120], fill="black", font=font)
            y += 22
            if y > 750:
                break
        img.save(out_path, "PNG")
        return [out_path]
    except Exception as e:
        log.error("Error al convertir Excel %s: %s", excel_file_path, e)
        raise ValueError(f"No se pudo procesar Excel: {excel_file_path}. Error: {e}")


def _render_texto_png(text_lines: list[str], base_name: str, output_folder: str) -> str:
    """Renderiza líneas de texto en un PNG 1200x800. Retorna la ruta generada."""
    if Image is None:
        raise ValueError("Pillow no disponible para renderizar texto a imagen")
    if not text_lines:
        text_lines = ["Documento sin contenido legible"]
    out_path = os.path.join(output_folder, f"{base_name}_txt_01.png")
    img = Image.new("RGB", (1200, 800), "white")
    draw = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype("arial.ttf", 14)
    except Exception:
        font = ImageFont.load_default()
    y = 10
    for line in text_lines[:35]:
        draw.text((10, y), line[:120], fill="black", font=font)
        y += 22
        if y > 750:
            break
    img.save(out_path, "PNG")
    return out_path


def word_to_imagenes(word_file_path: str, output_folder: str) -> list[str]:
    """Convierte un Word (.docx) a imagen para IA: extrae texto y genera PNG."""
    if not os.path.exists(word_file_path):
        raise FileNotFoundError(f"El archivo Word no existe: {word_file_path}")
    try:
        from docx import Document
    except ImportError as e:
        raise ValueError("python-docx no disponible para convertir Word a imagen") from e

    try:
        log.info("Procesando archivo Word: %s", word_file_path)
        doc = Document(word_file_path)
        text_lines = []
        for p in doc.paragraphs:
            if p.text and p.text.strip():
                text_lines.append(p.text.strip()[:200])
                if len(text_lines) > 50:
                    break
        for table in doc.tables:
            for row in table.rows:
                vals = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
                if vals:
                    text_lines.append(" | ".join(vals)[:200])
                if len(text_lines) > 60:
                    break
            if len(text_lines) > 60:
                break
        base_name = os.path.splitext(os.path.basename(word_file_path))[0]
        return [_render_texto_png(text_lines or ["Word sin contenido legible"], base_name, output_folder)]
    except Exception as e:
        log.error("Error al convertir Word %s: %s", word_file_path, e)
        raise ValueError(f"No se pudo procesar Word: {word_file_path}. Error: {e}")


def texto_to_imagenes(text_file_path: str, output_folder: str) -> list[str]:
    """Convierte un archivo de texto (.csv/.txt/.eml/.html/.msg) a imagen para IA."""
    if not os.path.exists(text_file_path):
        raise FileNotFoundError(f"El archivo de texto no existe: {text_file_path}")
    try:
        log.info("Procesando archivo de texto: %s", text_file_path)
        with open(text_file_path, "r", encoding="utf-8", errors="ignore") as f:
            contenido = f.read(20000)
        if text_file_path.lower().endswith((".html", ".htm")):
            contenido = re.sub(r"<script.*?</script>", " ", contenido, flags=re.DOTALL | re.IGNORECASE)
            contenido = re.sub(r"<style.*?</style>", " ", contenido, flags=re.DOTALL | re.IGNORECASE)
            contenido = re.sub(r"<[^>]+>", " ", contenido)
            contenido = re.sub(r"\s+", " ", contenido)
        text_lines = [ln.strip()[:200] for ln in contenido.splitlines() if ln.strip()][:60]
        base_name = os.path.splitext(os.path.basename(text_file_path))[0]
        return [_render_texto_png(text_lines or ["Texto sin contenido legible"], base_name, output_folder)]
    except Exception as e:
        log.error("Error al convertir texto %s: %s", text_file_path, e)
        raise ValueError(f"No se pudo procesar texto: {text_file_path}. Error: {e}")


TIPOS_EXCEL = ("xls", "xlsx", "xlsm", "xltx", "excel", "sheet")
TIPOS_WORD = ("docx", "docm")
TIPOS_TEXTO = ("csv", "txt", "eml", "html", "htm", "msg", "log")
TIPOS_IMAGEN_CONVERTIBLE = ("bmp", "tiff", "tif", "webp", "gif")


def archivo_a_imagenes(local_file: str, output_folder: str) -> list[str]:
    """Convierte un adjunto descargado a PNGs legibles por IA.

    Retorna lista vacía si el formato no es convertible (ej. ZIP, video).
    """
    ft = os.path.splitext(local_file)[1].lower().lstrip(".")
    if ft == "pdf":
        return pdf_to_imagenes(local_file, output_folder)
    if ft in TIPOS_EXCEL or "excel" in ft or "sheet" in ft:
        return excel_to_imagenes(local_file, output_folder)
    if ft in TIPOS_WORD:
        return word_to_imagenes(local_file, output_folder)
    if ft in TIPOS_TEXTO:
        return texto_to_imagenes(local_file, output_folder)
    if ft in ("jpg", "jpeg", "png"):
        return [local_file]
    if ft in TIPOS_IMAGEN_CONVERTIBLE and Image is not None:
        try:
            base_name = os.path.splitext(os.path.basename(local_file))[0]
            out_path = os.path.join(output_folder, f"{base_name}_conv_01.png")
            with Image.open(local_file) as img:
                img.convert("RGB").save(out_path, "PNG")
            return [out_path]
        except Exception as e:
            log.error("Error convirtiendo imagen %s: %s", local_file, e)
            return []
    return []


def AdjuntosSQIS(ID_CLAIM: str) -> dict[str, list[str]]:
    """Descarga los adjuntos asociados a un caso y retorna URLs SAS por documento."""
    download_path = str(ADJUNTOS_DIR / ID_CLAIM)

    if os.path.exists(download_path) and len(os.listdir(download_path)) > 0:
        log.info("Carpeta ya existe para %s, saltando descarga...", ID_CLAIM)
        return reconstruir_urls_existentes(download_path)

    os.makedirs(download_path, exist_ok=True)

    documentos_urls = URL_Adjuntos(ID_CLAIM)
    if documentos_urls is None:
        log.warning("No se encontraron adjuntos para %s", ID_CLAIM)
        return {}

    for item in documentos_urls:
        filename = item.get("filename")
        url = item.get("url")
        log.info("Descargando: %s", filename)

        doc_folder = os.path.splitext(filename)[0]
        local_doc_path = os.path.join(download_path, doc_folder)
        os.makedirs(local_doc_path, exist_ok=True)

        try:
            source_filename = os.path.basename(url.split("?")[0])
            base_name, file_extension = os.path.splitext(source_filename)
            base_name = base_name[:20]
            original_filename = f"{base_name}{file_extension}"

            local_file = os.path.join(local_doc_path, original_filename)

            r = requests.get(url)
            with open(local_file, "wb") as f:
                f.write(r.content)

            if original_filename.lower().endswith(".pdf"):
                try:
                    sanitized_file_path = sanitize_filename(local_file)
                    pdf_to_imagenes(os.path.normpath(sanitized_file_path), local_doc_path)
                    os.remove(local_file)
                except Exception as e:
                    log.error("Error al procesar PDF %s: %s", local_file, e)

        except Exception as e:
            log.error("Error al descargar %s: %s", url, e)

    urls_sas = {}
    for doc_folder in os.listdir(download_path):
        local_doc_path = os.path.join(download_path, doc_folder)
        if os.path.isdir(local_doc_path):
            urls_sas[doc_folder] = subir_imagenes_blob(f"{ID_CLAIM}/{doc_folder}", local_doc_path)

    return urls_sas


def sanitize_filename(file_path: str, max_length: int = 20) -> str:
    """Limpia el nombre del archivo de caracteres especiales y limita su longitud."""
    directory, filename = os.path.split(file_path)
    name, ext = os.path.splitext(filename)

    replacements = {"ñ": "n", "á": "a", "é": "e", "í": "i", "ó": "o", "ú": "u", "ü": "u"}
    for char, replacement in replacements.items():
        name = name.replace(char, replacement)

    sanitized_name = re.sub(r"[^a-zA-Z0-9_\-]", "", name.replace(" ", "_"))

    if not sanitized_name:
        sanitized_name = "archivo"

    sanitized_name = sanitized_name[:max_length]
    sanitized_filename = f"{sanitized_name}{ext}"
    sanitized_path = os.path.join(directory, sanitized_filename)

    if file_path != sanitized_path:
        shutil.move(file_path, sanitized_path)

    return sanitized_path
